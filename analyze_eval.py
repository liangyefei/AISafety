"""分析 3000 条评测的误判构成。"""
from __future__ import annotations

import sys
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook

wb = load_workbook("docs/评测报告_3000条.xlsx")
ws = wb["逐条明细"]
headers = [c.value for c in ws[1]]
rows = [dict(zip(headers, [c.value for c in r], strict=True)) for r in ws.iter_rows(min_row=2)]
wrong = [r for r in rows if r["是否正确"] == "✗"]

print(f"总判定正确 {len(rows) - len(wrong)}/{len(rows)}\n")
print("=== 误判类型构成 ===")
for key, count in Counter(f"{r['期望']} -> {r['label']}/{r['abuse_label']}" for r in wrong).most_common():
    print(f"  {key:26s} {count:4d}")

hard_miss = [r for r in wrong if r["期望"] == "spam" and r["label"] == "normal"]
print(f"\n=== 硬漏：期望 spam 判成 normal（{len(hard_miss)} 条）===")
for r in sorted(hard_miss, key=lambda x: -x["垃圾分"])[:10]:
    print(f'  垃圾分={r["垃圾分"]:.3f} 规则分={r["规则分"]:.3f}  {str(r["文本"])[:42]}')

hard_fp = [
    r for r in wrong
    if r["期望"] == "normal" and (r["label"] == "spam" or r["abuse_label"] == "spam")
]
print(f"\n=== 硬误杀：期望 normal 判成 spam（{len(hard_fp)} 条）===")
for r in sorted(hard_fp, key=lambda x: -max(x["垃圾分"], x["辱骂分"]))[:8]:
    print(f'  {r["label"]}/{r["abuse_label"]} 垃圾分={r["垃圾分"]:.3f} 辱骂分={r["辱骂分"]:.3f}  {str(r["文本"])[:38]}')

print("\n=== 分数分布：为什么 review 灰区多 ===")
buckets = Counter()
for r in rows:
    s = max(float(r["垃圾分"]), float(r["辱骂分"]))
    buckets[f"{int(s * 10) / 10:.1f}"] += 1
for key in sorted(buckets):
    print(f"  {key}  {buckets[key]:5d}")
